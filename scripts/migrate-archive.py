#!/usr/bin/env python3
"""
자료실 마이그레이션 — 구 사이트(아임웹) 엑셀 + 첨부파일 → woori API.

  python3 scripts/migrate-archive.py --dry-run --all   # 변환·매핑만 확인
  python3 scripts/migrate-archive.py --post 5466346    # 특정 글 1건 (파일럿)
  python3 scripts/migrate-archive.py --all             # 전체 (오래된 글부터)
  python3 scripts/migrate-archive.py --list-existing   # 게시판 현황
  python3 scripts/migrate-archive.py --delete IDX      # 특정 글 삭제

설치사례와 결정적으로 다른 점: **첨부파일이 엑셀에 없다.**
엑셀은 카테고리·제목·본문·날짜·조회수만 담고 있고, 정작 자료실의 본체인 PDF 는 구 사이트
상세페이지에만 있다. 그래서 게시물번호로 구 사이트를 긁어 파일을 받아온다.

날짜·조회수는 서버가 쓰기를 막으므로(date_active 는 생성시각 강제, count_read 는 조회 시에만
증가) extras 에 실어 보낸다. lib/contents.ts 의 parseExtras·toItem 이 그 값을 우선한다.
"""

import argparse
import datetime
import html as htmllib
import json
import os
import re
import sys
import unicodedata
import urllib.parse
from pathlib import Path

import openpyxl
import requests

ROOT = Path(__file__).resolve().parent.parent
LEDGER = ROOT / "scripts" / ".migrate-archive-ledger.json"

BOARD = "archive"
MODULE_IDX = "019cb2f7-0d89-75b4-a16b-6019e8ee9782"
FILE_TYPE_GENERAL = 1  # 첨부파일 (다운로드용)
FILE_TYPE_IMAGE = 3    # 본문 이미지
PERMIT_PUBLIC = 0
DATE_END_FAR = "2099-12-31"

# 구 사이트 — 아직 살아 있다. 게시물번호가 그대로 idx 다.
OLD_BASE = "https://www.arumtech.co.kr"
OLD_Q = "YToxOntzOjEyOiJrZXl3b3JkX3R5cGUiO3M6MzoiYWxsIjt9"
OLD_UA = {
    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120 Safari/537.36"
}

# 엑셀 파일명 → 사이트 카테고리 (lib/contents.ts:30 의 BOARD_CATEGORIES.archive).
# 엑셀 A열 카테고리는 쓰지 않는다 — '도면'·'새 카테고리'·공란이 섞여 있어 쓸모가 없다.
# 매뉴얼→메뉴얼, 카다로그→카탈로그는 표기가 다르다. 사이트 값이 정본이다.
CATEGORY_BY_FILE = {
    "기술자료": "기술자료",
    "도면자료": "도면자료",
    "매뉴얼": "메뉴얼",
    "물가정보": "물가정보",
    "시방서": "시방서",
    "카다로그": "카탈로그",
}

# 2026년에 외부인이 올린 카지노 홍보글. 첨부도 없다.
SPAM_POST_NOS = {"170593238", "170653024"}


def load_env() -> dict:
    path = ROOT / ".env.local"
    if not path.exists():
        sys.exit("[중단] .env.local 이 없습니다.")
    env = {}
    for line in path.read_text().splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            env[k.strip()] = v.strip()
    return env


ENV = load_env()
API = ENV.get("NEXT_PUBLIC_WOORI_API_URL", "")
SITE = ENV.get("NEXT_PUBLIC_WOORI_SITE_ID", "")
FILE_URL = ENV.get("NEXT_PUBLIC_WOORI_FILE_URL", "")


def login() -> str:
    uid, pw = ENV.get("WOORI_ADMIN_ID"), ENV.get("WOORI_ADMIN_PW")
    if not uid or not pw:
        sys.exit("[중단] .env.local 에 WOORI_ADMIN_ID / WOORI_ADMIN_PW 를 넣어주세요.")
    r = requests.post(
        f"{API}/members/login",
        headers={"Content-Type": "application/json", "x-site": SITE},
        json={"member_id": uid, "member_pw": pw, "member_type": 0},
        timeout=30,
    )
    body = r.json()
    token = (body.get("auth") or {}).get("access_token")
    if not token:
        sys.exit(f"[중단] 로그인 실패: {body.get('message')}")
    print(f"  로그인 성공 — {(body.get('data') or {}).get('member_id')}")
    return token


def read_rows() -> list[dict]:
    """엑셀 6개를 읽는다. 10컬럼·11컬럼 두 스키마가 섞여 있어 헤더명으로 분기한다."""
    rows = []
    resource = ROOT / "resource"
    for f in sorted(os.listdir(resource)):
        if not f.endswith(".xlsx"):
            continue
        # macOS 는 한글 파일명을 NFD 로 저장해 그대로 비교하면 빗나간다.
        name = unicodedata.normalize("NFC", f)
        stem = name.split("_")[0]
        if stem not in CATEGORY_BY_FILE:
            continue  # 설치사례 등 대상 외
        ws = openpyxl.load_workbook(resource / f)["Sheet1"]
        hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col = {h: i + 1 for i, h in enumerate(hdr) if h}
        for need in ("게시물 번호", "제목", "내용", "작성시각", "조회수"):
            if need not in col:
                sys.exit(f"[중단] {name}: '{need}' 컬럼이 없습니다. 헤더={hdr}")
        for r in range(2, ws.max_row + 1):
            post_no = ws.cell(r, col["게시물 번호"]).value
            if not post_no:
                continue
            post_no = str(post_no).strip()
            if post_no in SPAM_POST_NOS:
                continue
            rows.append({
                "post_no": post_no,
                "category": CATEGORY_BY_FILE[stem],
                "source": stem,
                "title": str(ws.cell(r, col["제목"]).value or "").strip(),
                "html": ws.cell(r, col["내용"]).value or "",
                "written_at": str(ws.cell(r, col["작성시각"]).value or ""),
                "ip": str(ws.cell(r, col.get("IP 주소", 0)).value or "0.0.0.0") if "IP 주소" in col else "0.0.0.0",
                "views": int(ws.cell(r, col["조회수"]).value or 0),
            })
    return rows


def order_val_of(written_at: str) -> int:
    dt = datetime.datetime.strptime(written_at, "%Y-%m-%d %H:%M:%S")
    kst = datetime.timezone(datetime.timedelta(hours=9))
    return int(dt.replace(tzinfo=kst).timestamp() * 1000)


# ── 구 사이트에서 첨부 받아오기 ────────────────────────────────────────────────

def fetch_old_attachments(post_no: str) -> list[tuple[str, bytes]]:
    """
    구 사이트 상세페이지를 열고 첨부파일을 받아 [(파일명, 바이트)] 로 돌려준다.

    다운로드 링크의 base64 안에는 post_download_token 이 들어 있고 페이지를 열 때마다 새로
    발급된다. 반드시 **같은 세션에서 페이지를 받은 직후** 내려받아야 한다 — 링크를 모아뒀다가
    나중에 쓰면 404 가 난다.
    """
    s = requests.Session()
    s.headers.update(OLD_UA)
    page = s.get(f"{OLD_BASE}/?q={OLD_Q}&bmode=view&idx={post_no}&t=board", timeout=30)
    page.raise_for_status()

    links = re.findall(r'href="(/post_file_download\.cm\?c=[^"]+)"', page.text)
    out = []
    for link in links:
        r = s.get(OLD_BASE + htmllib.unescape(link), headers={"Referer": page.url}, timeout=300)
        r.raise_for_status()
        cd = r.headers.get("Content-Disposition", "")
        m = re.search(r"filename\*?=(?:UTF-8'')?\"?([^\";]+)", cd)
        fname = urllib.parse.unquote(m.group(1)).strip() if m else f"{post_no}.bin"
        out.append((fname, r.content))
    return out


# ── 우리 파일서버 업로드 ──────────────────────────────────────────────────────

def storage_put(fname: str, blob: bytes, file_type: int) -> dict:
    """
    post.php 에 올리고 메타를 돌려받는다.

    헤더를 하나도 붙이지 않는다 — 파일서버(iwinv)는 Content-Type·Authorization 외의 헤더를
    CORS 로 막는다. post.php 는 인증도 x-site 도 보지 않는다 (lib/contents-write.ts:203-209).
    """
    r = requests.post(
        f"{FILE_URL}/files/post.php",
        files={"file": (fname, blob)},
        data={
            "file_path": BOARD,
            "file_type": str(file_type),
            "permit_level": str(PERMIT_PUBLIC),
        },
        timeout=600,
    )
    r.raise_for_status()
    meta = (r.json() or {}).get("data") or {}
    if not meta.get("file_name_real"):
        raise RuntimeError(f"업로드 응답에 file_name_real 없음: {r.text[:200]}")
    return meta


def public_url(meta: dict) -> str:
    """lib/contents-write.ts:258-271 의 publicUrl 과 동일 규칙. 실제 경로는 403 이라 get.php 로만 받는다."""
    q = urllib.parse.urlencode({
        "permit_level": str(PERMIT_PUBLIC),
        "file_path": meta.get("file_path") or "",
        "file_name_origin": meta.get("file_name_origin") or "",
        "file_name_real": meta.get("file_name_real") or "",
    })
    return f"{FILE_URL}/files/get.php?{q}"


def register_file(meta: dict, token: str) -> str:
    """
    첨부파일은 post.php 업로드만으론 부족하다. woori API 에 메타를 등록해 idx 를 받아야
    게시글에 붙일 수 있다 (본문 이미지는 이 단계가 없다 — URL 만 쓰면 된다).
    """
    r = requests.post(
        f"{API}/files",
        headers={
            "Content-Type": "application/json",
            "x-site": SITE,
            "Authorization": f"Bearer {token}",
        },
        json={k: meta.get(k) for k in (
            "file_ext", "file_name_origin", "file_name_real",
            "file_path", "file_size", "file_type", "permit_level",
        )},
        timeout=60,
    )
    idx = ((r.json() or {}).get("data") or {}).get("idx")
    if not idx:
        raise RuntimeError(f"파일 등록 실패: {r.text[:200]}")
    return idx


def rewrite_body(body: str, dry: bool) -> tuple[str, int]:
    """본문의 구 CDN 이미지를 우리 파일서버로 옮기고 src 를 치환한다."""
    srcs = re.findall(r'<img\b[^>]*?\bsrc="([^"]+)"', body, re.I)
    n = 0
    for src in dict.fromkeys(srcs):
        n += 1
        if dry:
            continue
        got = requests.get(src, headers={**OLD_UA, "Referer": f"{OLD_BASE}/"}, timeout=120)
        got.raise_for_status()
        name = os.path.basename(urllib.parse.urlparse(src).path) or "image.jpg"
        url = public_url(storage_put(name, got.content, FILE_TYPE_IMAGE))
        body = body.replace(f'src="{src}"', f'src="{htmllib.escape(url, quote=True)}"')
    return body, n


def create_post(item: dict, body: str, file_idxs: list[str], token: str) -> str:
    payload = {
        "module_idx": MODULE_IDX,
        "title": item["title"],
        "content": body,
        "is_html": 1,
        # date·views 는 서버가 쓰기를 막는 date_active·count_read 대신 쓰는 값이다.
        # listImage 는 비운다 — 자료실 목록은 썸네일 대신 파일 아이콘만 쓴다 (BoardForm.tsx:61).
        "extras": json.dumps({
            "isPinned": 0,
            "category": item["category"],
            "listImage": "",
            "date": item["written_at"][:10],
            "views": item["views"],
        }, ensure_ascii=False),
        "ip": item["ip"],
        "order_val": order_val_of(item["written_at"]),
        "date_end": DATE_END_FAR,
        "is_hidden_list": 0,
        **({"file_idx": file_idxs} if file_idxs else {}),
    }
    r = requests.post(
        f"{API}/contents",
        headers={
            "Content-Type": "application/json",
            "x-site": SITE,
            "Authorization": f"Bearer {token}",
        },
        json=payload,
        timeout=120,
    )
    body_json = r.json()
    if body_json.get("statusCode") != 200:
        raise RuntimeError(f"등록 실패: {body_json.get('message')}")
    idx = (body_json.get("data") or {}).get("idx")
    if not idx:
        raise RuntimeError(f"등록 응답에 idx 없음: {str(body_json)[:200]}")
    return idx


def delete_post(idx: str, token: str) -> None:
    r = requests.delete(
        f"{API}/contents",
        headers={"Content-Type": "application/json", "x-site": SITE,
                 "Authorization": f"Bearer {token}"},
        json={"idx": idx}, timeout=30,
    )
    if r.json().get("statusCode") != 200:
        raise RuntimeError(f"삭제 실패: {r.json().get('message')}")


def list_existing() -> list[dict]:
    # page_size 가 맞는 이름이다. limit/page 를 보내면 API 가 무시하고 20건만 준다.
    r = requests.get(
        f"{API}/contents/list", headers={"x-site": SITE},
        params={"module_idx": MODULE_IDX, "page_size": 300, "content_load": 1, "file_load": 1},
        timeout=90,
    )
    return r.json().get("list") or []


def load_ledger() -> dict:
    return json.loads(LEDGER.read_text()) if LEDGER.exists() else {}


def save_ledger(led: dict) -> None:
    LEDGER.write_text(json.dumps(led, ensure_ascii=False, indent=2))


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--post", help="특정 게시물번호 1건")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--dry-run", action="store_true", help="다운로드·업로드·등록 없이 변환만")
    ap.add_argument("--list-existing", action="store_true")
    ap.add_argument("--delete", metavar="IDX")
    a = ap.parse_args()

    if a.list_existing:
        items = list_existing()
        print(f"자료실: {len(items)}건")
        for i in items:
            ex = json.loads(i.get("extras") or "{}")
            print(f"  {i['idx']}  {(i.get('title') or '')[:30]:32s} "
                  f"cat={ex.get('category', '')!r:10s} date={ex.get('date', '-')} "
                  f"views={ex.get('views', i.get('count_read'))} files={len(i.get('files') or [])}")
        return

    if a.delete:
        delete_post(a.delete, login())
        print(f"삭제 완료: {a.delete}")
        return

    rows = read_rows()
    if a.post:
        rows = [x for x in rows if x["post_no"] == a.post]
        if not rows:
            sys.exit(f"[중단] 게시물번호 {a.post} 을 찾지 못했습니다.")
    elif not a.all:
        sys.exit("--post N / --all 중 하나를 지정하세요.")

    rows.sort(key=lambda x: x["written_at"])  # 오래된 글부터
    ledger = load_ledger()
    token = None if a.dry_run else login()
    done = skipped = failed = 0

    for item in rows:
        tag = f"[{item['category']:6s}] {item['post_no']:>9s} {item['title'][:30]}"
        if item["post_no"] in ledger:
            print(f"  건너뜀 {tag} — 이미 등록됨")
            skipped += 1
            continue
        print(f"\n  {tag}")
        try:
            if a.dry_run:
                _, imgs = rewrite_body(item["html"], True)
                print(f"      본문 {len(item['html'])}자, 이미지 {imgs}장 | "
                      f"date={item['written_at'][:10]} views={item['views']} "
                      f"order_val={order_val_of(item['written_at'])}")
                done += 1
                continue

            atts = fetch_old_attachments(item["post_no"])
            file_idxs = []
            for fname, blob in atts:
                meta = storage_put(fname, blob, FILE_TYPE_GENERAL)
                file_idxs.append(register_file(meta, token))
                print(f"      첨부 {len(blob)/1024/1024:6.1f}MB  {fname}")
            if not atts:
                print("      첨부 없음 (본문만)")

            body, imgs = rewrite_body(item["html"], False)
            if imgs:
                print(f"      본문 이미지 {imgs}장 이전")

            idx = create_post(item, body, file_idxs, token)
            ledger[item["post_no"]] = {
                "idx": idx, "title": item["title"], "category": item["category"],
                "files": len(file_idxs),
            }
            save_ledger(ledger)
            print(f"      등록 완료 → {idx}")
            done += 1
        except Exception as e:
            print(f"      [실패] {type(e).__name__}: {e}")
            failed += 1

    print(f"\n완료: {done}건 / 건너뜀 {skipped}건 / 실패 {failed}건")
    if failed:
        sys.exit(1)


if __name__ == "__main__":
    main()
