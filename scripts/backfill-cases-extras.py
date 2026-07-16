#!/usr/bin/env python3
"""
설치사례 100건에 원본 게시일·조회수를 채워 넣는다 (이미 등록된 글을 수정).

  python3 scripts/backfill-cases-extras.py --dry-run
  python3 scripts/backfill-cases-extras.py --apply

설치사례를 옮길 때는 날짜·조회수를 포기했다가, 자료실을 옮기며 extras 로 복구하기로 하면서
두 게시판의 표시가 어긋나게 됐다. 이 스크립트가 그걸 맞춘다.

글을 다시 만들지 않는다 — PUT 으로 extras 만 덧씌운다. 본문 이미지 197장(96MB)을 다시
올릴 이유가 없다. 어느 글이 어느 엑셀 행인지는 마이그레이션 대장이 알고 있다.
"""

import argparse
import datetime
import json
import os
import sys
import unicodedata
from pathlib import Path

import openpyxl
import requests

ROOT = Path(__file__).resolve().parent.parent
LEDGER = ROOT / "scripts" / ".migrate-cases-ledger.json"


def load_env() -> dict:
    env = {}
    for line in (ROOT / ".env.local").read_text().splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            env[k.strip()] = v.strip()
    return env


ENV = load_env()
API = ENV["NEXT_PUBLIC_WOORI_API_URL"]
SITE = ENV["NEXT_PUBLIC_WOORI_SITE_ID"]


def login() -> str:
    r = requests.post(
        f"{API}/members/login",
        headers={"Content-Type": "application/json", "x-site": SITE},
        json={"member_id": ENV["WOORI_ADMIN_ID"], "member_pw": ENV["WOORI_ADMIN_PW"],
              "member_type": 0},
        timeout=30,
    ).json()
    tok = (r.get("auth") or {}).get("access_token")
    if not tok:
        sys.exit(f"[중단] 로그인 실패: {r.get('message')}")
    return tok


def excel_rows() -> dict:
    """게시물번호 → {date, views}. 설치사례 엑셀에서 읽는다."""
    resource = ROOT / "resource"
    src = None
    for f in os.listdir(resource):
        # macOS 는 한글 파일명을 NFD 로 저장한다 — 정규화해서 비교한다.
        if f.endswith(".xlsx") and "설치사례" in unicodedata.normalize("NFC", f):
            src = resource / f
    if not src:
        sys.exit("[중단] resource/ 에서 설치사례 엑셀을 찾지 못했습니다.")
    ws = openpyxl.load_workbook(src)["Sheet1"]
    hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    col = {h: i + 1 for i, h in enumerate(hdr) if h}
    out = {}
    for r in range(2, ws.max_row + 1):
        pid = ws.cell(r, col["게시물 번호"]).value
        if not pid:
            continue
        out[str(pid).strip()] = {
            "date": str(ws.cell(r, col["작성시각"]).value or "")[:10],
            "views": int(ws.cell(r, col["조회수"]).value or 0),
        }
    return out


def fetch_extras(idx: str) -> str:
    """현재 extras 를 그대로 가져온다. 덮어쓰면 카테고리·목록이미지가 날아간다."""
    d = requests.get(f"{API}/contents", headers={"x-site": SITE},
                     params={"idx": idx}, timeout=30).json()
    return ((d.get("data") or d) or {}).get("extras") or "{}"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="실제로 수정")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    if not (a.apply or a.dry_run):
        sys.exit("--dry-run 또는 --apply 를 지정하세요.")

    if not LEDGER.exists():
        sys.exit("[중단] 설치사례 대장이 없습니다 — 어느 글이 어느 행인지 알 수 없습니다.")
    ledger = json.loads(LEDGER.read_text())
    rows = excel_rows()
    token = None if a.dry_run else login()

    done = failed = 0
    for post_no, rec in sorted(ledger.items(), key=lambda x: x[1]["row"]):
        idx = rec["idx"]
        src = rows.get(post_no)
        if not src:
            print(f"  [건너뜀] {post_no} — 엑셀에 없음")
            continue

        ex = json.loads(fetch_extras(idx))
        merged = {**ex, "date": src["date"], "views": src["views"]}
        label = f"{rec['title'][:26]:28s} {src['date']} 조회 {src['views']:>5d}"
        if a.dry_run:
            print(f"  {label}  cat={ex.get('category')!r}")
            done += 1
            continue
        try:
            r = requests.put(
                f"{API}/contents",
                headers={"Content-Type": "application/json", "x-site": SITE,
                         "Authorization": f"Bearer {token}"},
                json={"idx": idx, "extras": json.dumps(merged, ensure_ascii=False)},
                timeout=60,
            ).json()
            if r.get("statusCode") != 200:
                raise RuntimeError(r.get("message"))
            print(f"  O {label}")
            done += 1
        except Exception as e:
            print(f"  X {label} — {e}")
            failed += 1

    print(f"\n완료: {done}건 / 실패 {failed}건")
    if failed:
        sys.exit(1)


if __name__ == "__main__":
    main()
