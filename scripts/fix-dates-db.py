#!/usr/bin/env python3
"""
이관 글의 date_active(게시일)·count_read(조회수)를 PostgreSQL 에서 직접 보정한다.

  python3 scripts/fix-dates-db.py --dry-run   # SQL 을 만들고 대상만 SELECT (수정 없음)
  python3 scripts/fix-dates-db.py --apply      # 실제 UPDATE

woori API 는 date_active·count_read 의 쓰기를 막는다(전자는 생성시각 강제, 후자는 조회 시
증가만). API 로는 원본값을 넣을 수 없어 extras 로 우회했었는데, 이 스크립트가 DB 를 직접
고쳐 근본적으로 해결한다. 동시에 extras 에서 우회용으로 넣었던 date·views 키를 제거한다
(isPinned·category·listImage 는 보존).

접속: SSH(DEPLOY_SSH_*)로 서버에 들어가 그 안의 psql 로 붙는다. postgres 는 127.0.0.1:5432
에만 떠 있어 외부에서 직접 못 붙는다. SQL 은 stdin 으로 넘겨 따옴표 문제를 피한다.
대상은 대장(ledger) + 엑셀에서 만든다. 대장에 없는 글·다른 게시판·새 글은 건드리지 않는다.
"""

import argparse
import json
import os
import subprocess
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
ARCHIVE_CATS = {"기술자료", "도면자료", "매뉴얼", "물가정보", "시방서", "카다로그"}


def env() -> dict:
    e = {}
    for line in (ROOT / ".env.local").read_text().splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            e[k.strip()] = v.strip()
    return e


ENV = env()


def excel_map(is_cases: bool) -> dict:
    res = {}
    for f in os.listdir(ROOT / "resource"):
        if not f.endswith(".xlsx"):
            continue
        n = unicodedata.normalize("NFC", f)
        if is_cases and "설치사례" not in n:
            continue
        if not is_cases and ("설치사례" in n or n.split("_")[0] not in ARCHIVE_CATS):
            continue
        ws = openpyxl.load_workbook(ROOT / "resource" / f)["Sheet1"]
        hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col = {h: i + 1 for i, h in enumerate(hdr) if h}
        for r in range(2, ws.max_row + 1):
            pid = ws.cell(r, col["게시물 번호"]).value
            if not pid:
                continue
            res[str(pid).strip()] = {
                "written_at": str(ws.cell(r, col["작성시각"]).value or ""),
                "views": int(ws.cell(r, col["조회수"]).value or 0),
            }
    return res


def targets() -> list[dict]:
    rows = []
    for ledger, is_cases in [
        ("scripts/.migrate-cases-ledger.json", True),
        ("scripts/.migrate-archive-ledger.json", False),
    ]:
        path = ROOT / ledger
        if not path.exists():
            sys.exit(f"[중단] 대장 없음: {ledger}")
        led = json.loads(path.read_text())
        xl = excel_map(is_cases)
        for pid, rec in led.items():
            src = xl.get(pid)
            if not src:
                sys.exit(f"[중단] {pid} 이 엑셀에 없습니다.")
            rows.append({"idx": rec["idx"], "written_at": src["written_at"], "views": src["views"]})
    return rows


def build_sql(rows: list[dict], apply: bool) -> str:
    """
    임시 테이블에 (idx, 게시일, 조회수) 를 실어 한 번에 UPDATE 한다. 행마다 UPDATE 를 날리면
    237번 왕복이라 느리다. 값은 psql 파라미터가 아니라 리터럴로 박되, 형식이 고정된 값
    (UUID·'YYYY-MM-DD HH:MM:SS'·정수)만 쓰므로 인젝션 여지가 없다.

    date_active 는 작성시각을 KST(+09) 로 명시해 넣는다. 화면은 이 값을 slice(0,10) 하므로
    타임존만 맞으면 날짜가 안 밀린다.
    extras 는 text 라 jsonb 로 캐스팅해 date·views 키를 지운 뒤 다시 text 로 되돌린다.
    """
    values = ",\n".join(
        f"('{r['idx']}'::uuid, '{r['written_at']}+09'::timestamptz, {int(r['views'])})"
        for r in rows
    )
    sql = f"""\
BEGIN;

CREATE TEMP TABLE _fix (idx uuid, d timestamptz, v integer) ON COMMIT DROP;
INSERT INTO _fix (idx, d, v) VALUES
{values};

-- 안전장치: 대상이 _fix 에 있는 행만 고친다.
UPDATE wr_contents_t c
SET date_active = f.d,
    count_read  = f.v,
    extras      = ((c.extras::jsonb) - 'date' - 'views')::text
FROM _fix f
WHERE c.idx = f.idx;

-- 결과 확인 (dry-run 은 여기 SELECT 만 보고 롤백).
SELECT count(*) AS updated FROM wr_contents_t c JOIN _fix f ON c.idx = f.idx
  WHERE c.date_active = f.d AND c.count_read = f.v;

SELECT c.title,
       to_char(c.date_active AT TIME ZONE 'Asia/Seoul', 'YYYY-MM-DD') AS shown_date,
       c.count_read,
       c.extras
FROM wr_contents_t c JOIN _fix f ON c.idx = f.idx
ORDER BY c.date_active LIMIT 4;

{'COMMIT;' if apply else 'ROLLBACK;'}
"""
    return sql


def psql(sql: str) -> str:
    """SSH 로 서버에 들어가 psql 에 SQL 을 stdin 으로 넘긴다."""
    host = ENV["DEPLOY_SSH_HOST"]
    user = ENV["DEPLOY_SSH_USER"]
    dbpw = ENV["WOORI_DB_PASSWORD"]
    remote = (
        f"PGPASSWORD='{dbpw}' /usr/pgsql-13/bin/psql "
        f"-h {ENV['WOORI_DB_HOST']} -p {ENV['WOORI_DB_PORT']} "
        f"-U {ENV['WOORI_DB_USER']} -d {ENV['WOORI_DB_NAME']} "
        f"-v ON_ERROR_STOP=1 -f -"
    )
    p = subprocess.run(
        ["sshpass", "-e", "ssh", "-o", "StrictHostKeyChecking=accept-new",
         f"{user}@{host}", remote],
        input=sql, capture_output=True, text=True,
        env={**os.environ, "SSHPASS": ENV["DEPLOY_SSH_PASS"]},
        timeout=300,
    )
    out = "\n".join(l for l in (p.stdout + p.stderr).splitlines()
                    if not any(x in l for x in ("post-quantum", "store now", "openssh.com", "WARNING")))
    if p.returncode != 0:
        raise RuntimeError(f"psql 실패:\n{out}")
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    if not (a.apply or a.dry_run):
        sys.exit("--dry-run 또는 --apply 를 지정하세요.")

    rows = targets()
    print(f"대상 {len(rows)}건 (설치사례 100 + 자료실 137 = 237 기대)")
    sql = build_sql(rows, apply=a.apply)

    mode = "실제 UPDATE (COMMIT)" if a.apply else "DRY-RUN (ROLLBACK — 아무것도 안 바뀜)"
    print(f"모드: {mode}\n")
    print(psql(sql))


if __name__ == "__main__":
    main()
